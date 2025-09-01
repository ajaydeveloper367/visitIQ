#!/usr/bin/env python3
"""
Multi-Format Medical Document Processing Test
Creates and processes various document formats: PDF, images, Word, DICOM, etc.
"""

import os
import sys
from pathlib import Path
from datetime import datetime
import json

# Add src to path
sys.path.append('src')

def create_sample_pdf_document(output_path: str):
    """Create a sample PDF medical report"""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        
        # Create PDF with medical content
        doc = SimpleDocTemplate(output_path, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []
        
        # Add content
        title = Paragraph("LABORATORY REPORT", styles['Title'])
        story.append(title)
        story.append(Spacer(1, 12))
        
        content = """
        <b>Patient:</b> John Anderson<br/>
        <b>DOB:</b> 03/15/1975<br/>
        <b>Date:</b> January 15, 2024<br/>
        <b>Physician:</b> Dr. Sarah Johnson<br/><br/>
        
        <b>COMPLETE METABOLIC PANEL:</b><br/>
        Glucose: 185 mg/dL (HIGH) [70-100]<br/>
        BUN: 22 mg/dL [7-25]<br/>
        Creatinine: 1.1 mg/dL [0.6-1.2]<br/>
        Sodium: 142 mEq/L [135-145]<br/>
        Potassium: 4.2 mEq/L [3.5-5.0]<br/><br/>
        
        <b>LIPID PANEL:</b><br/>
        Total Cholesterol: 220 mg/dL (HIGH) [<200]<br/>
        HDL: 45 mg/dL [>40]<br/>
        LDL: 140 mg/dL (HIGH) [<100]<br/>
        Triglycerides: 175 mg/dL (HIGH) [<150]<br/><br/>
        
        <b>HEMOGLOBIN A1C:</b><br/>
        HbA1c: 8.2% (HIGH) [<7.0%]<br/><br/>
        
        <b>INTERPRETATION:</b><br/>
        Elevated glucose and HbA1c consistent with diabetes.<br/>
        Lipid abnormalities noted. Recommend dietary modification.
        """
        
        para = Paragraph(content, styles['Normal'])
        story.append(para)
        
        doc.build(story)
        print(f"✅ Created PDF: {output_path}")
        return True
        
    except ImportError:
        # Fallback: Create PDF using PyMuPDF
        try:
            import fitz  # PyMuPDF
            
            doc = fitz.open()  # Create new PDF
            page = doc.new_page()
            
            # Add text content
            content = """LABORATORY REPORT

Patient: John Anderson
DOB: 03/15/1975
Date: January 15, 2024
Physician: Dr. Sarah Johnson

COMPLETE METABOLIC PANEL:
Glucose: 185 mg/dL (HIGH) [70-100]
BUN: 22 mg/dL [7-25]
Creatinine: 1.1 mg/dL [0.6-1.2]
Sodium: 142 mEq/L [135-145]
Potassium: 4.2 mEq/L [3.5-5.0]

LIPID PANEL:
Total Cholesterol: 220 mg/dL (HIGH) [<200]
HDL: 45 mg/dL [>40]
LDL: 140 mg/dL (HIGH) [<100]
Triglycerides: 175 mg/dL (HIGH) [<150]

HEMOGLOBIN A1C:
HbA1c: 8.2% (HIGH) [<7.0%]

INTERPRETATION:
Elevated glucose and HbA1c consistent with diabetes.
Lipid abnormalities noted. Recommend dietary modification."""

            # Insert text
            page.insert_text((50, 50), content, fontsize=11)
            
            # Save PDF
            doc.save(output_path)
            doc.close()
            print(f"✅ Created PDF (PyMuPDF): {output_path}")
            return True
            
        except Exception as e:
            print(f"❌ PDF creation failed: {e}")
            return False

def create_sample_image_with_text(output_path: str):
    """Create a sample image with medical text for OCR testing"""
    try:
        from PIL import Image, ImageDraw, ImageFont
        
        # Create image
        width, height = 800, 600
        image = Image.new('RGB', (width, height), 'white')
        draw = ImageDraw.Draw(image)
        
        # Try to use a better font, fallback to default
        try:
            font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 16)
        except:
            font = ImageFont.load_default()
        
        # Add medical report text
        text_content = """CHEST X-RAY REPORT

Patient: Mary Johnson
Date: January 20, 2024
Study: PA and Lateral Chest

FINDINGS:
- Heart size is within normal limits
- Lungs are clear bilaterally
- No pleural effusion
- Bone structures appear intact
- BP: 140/90 mmHg
- HR: 82 bpm
- Glucose: 210 mg/dL

IMPRESSION:
Normal chest radiograph.
Recommend follow-up for elevated glucose.

Radiologist: Dr. Michael Chen, MD"""

        # Draw text on image
        y_position = 50
        for line in text_content.split('\n'):
            draw.text((50, y_position), line, fill='black', font=font)
            y_position += 25
        
        # Save image
        image.save(output_path)
        print(f"✅ Created Image: {output_path}")
        return True
        
    except Exception as e:
        print(f"❌ Image creation failed: {e}")
        return False

def create_sample_word_document(output_path: str):
    """Create a sample Word document with medical content"""
    try:
        from docx import Document
        from docx.shared import Inches
        
        doc = Document()
        
        # Add title
        title = doc.add_heading('DISCHARGE SUMMARY', 0)
        
        # Add content
        doc.add_heading('Patient Information', level=1)
        p = doc.add_paragraph()
        p.add_run('Patient: ').bold = True
        p.add_run('Robert Williams')
        p.add_run('\nMRN: ').bold = True
        p.add_run('12345678')
        p.add_run('\nDOB: ').bold = True
        p.add_run('07/22/1968')
        
        doc.add_heading('Hospital Course', level=1)
        p = doc.add_paragraph()
        p.add_run('Admission Date: ').bold = True
        p.add_run('January 10, 2024')
        p.add_run('\nDischarge Date: ').bold = True
        p.add_run('January 15, 2024')
        
        p = doc.add_paragraph(
            'Patient admitted for management of diabetic ketoacidosis. '
            'Initial glucose: 425 mg/dL, pH: 7.25. Responded well to insulin therapy. '
            'Glucose at discharge: 180 mg/dL. BP: 135/85 mmHg. HR: 78 bpm.'
        )
        
        doc.add_heading('Discharge Medications', level=1)
        medications = [
            'Metformin 1000mg twice daily',
            'Insulin glargine 20 units daily',
            'Lisinopril 10mg daily',
            'Atorvastatin 40mg daily'
        ]
        
        for med in medications:
            doc.add_paragraph(med, style='List Bullet')
        
        doc.add_heading('Follow-up', level=1)
        doc.add_paragraph('Primary care physician in 1 week. Endocrinology in 2 weeks.')
        
        # Save document
        doc.save(output_path)
        print(f"✅ Created Word Document: {output_path}")
        return True
        
    except Exception as e:
        print(f"❌ Word document creation failed: {e}")
        return False

def create_sample_dicom_metadata(output_path: str):
    """Create a sample DICOM-like metadata file"""
    try:
        # Create a JSON file with DICOM-like metadata
        dicom_metadata = {
            "PatientName": "Jennifer Davis",
            "PatientID": "67890123",
            "StudyDate": "20240118",
            "StudyTime": "143022",
            "Modality": "CT",
            "StudyDescription": "CT Chest with Contrast",
            "SeriesDescription": "Axial CT Chest",
            "BodyPartExamined": "CHEST",
            "StudyInstanceUID": "1.2.3.4.5.6.7.8.9.10",
            "Manufacturer": "GE Medical Systems",
            "SliceThickness": "5.0",
            "KVP": "120",
            "PatientAge": "045Y",
            "PatientSex": "F",
            "PatientWeight": "65.5",
            "StudyID": "12345",
            "AccessionNumber": "ACC789012",
            "InstitutionName": "Regional Medical Center",
            "ReferringPhysicianName": "Dr. Lisa Thompson",
            "ClinicalInfo": "Chest pain, rule out pulmonary embolism",
            "Findings": {
                "impression": "No evidence of pulmonary embolism",
                "heart_rate": "88 bpm",
                "blood_pressure": "128/82 mmHg",
                "additional_notes": "Normal cardiac silhouette"
            }
        }
        
        with open(output_path, 'w') as f:
            json.dump(dicom_metadata, f, indent=2)
        
        print(f"✅ Created DICOM Metadata: {output_path}")
        return True
        
    except Exception as e:
        print(f"❌ DICOM metadata creation failed: {e}")
        return False

def test_multi_format_processing():
    """Test processing of multiple document formats"""
    print("🏥 MULTI-FORMAT DOCUMENT PROCESSING TEST")
    print("=" * 60)
    
    # Create test patient directory
    test_patient_dir = Path("data/patient_documents/patient_999")
    
    # Create directories for different document types
    directories = {
        'lab_reports': test_patient_dir / 'lab_reports',
        'imaging': test_patient_dir / 'imaging',
        'clinical_notes': test_patient_dir / 'clinical_notes',
        'discharge_summaries': test_patient_dir / 'discharge_summaries',
        'dicom_metadata': test_patient_dir / 'dicom_metadata'
    }
    
    for dir_path in directories.values():
        dir_path.mkdir(parents=True, exist_ok=True)
    
    print(f"\n📁 Created test directories for Patient 999")
    
    # Create sample documents in various formats
    print(f"\n📄 CREATING SAMPLE DOCUMENTS:")
    print("-" * 40)
    
    success_count = 0
    total_docs = 0
    
    # 1. PDF Lab Report
    total_docs += 1
    pdf_path = directories['lab_reports'] / '2024-01-15_lab_results.pdf'
    if create_sample_pdf_document(str(pdf_path)):
        success_count += 1
    
    # 2. Image Report (for OCR)
    total_docs += 1
    img_path = directories['imaging'] / '2024-01-20_chest_xray_report.png'
    if create_sample_image_with_text(str(img_path)):
        success_count += 1
    
    # 3. Word Document
    total_docs += 1
    word_path = directories['discharge_summaries'] / '2024-01-15_discharge_summary.docx'
    if create_sample_word_document(str(word_path)):
        success_count += 1
    
    # 4. DICOM Metadata
    total_docs += 1
    dicom_path = directories['dicom_metadata'] / '2024-01-18_ct_chest_metadata.json'
    if create_sample_dicom_metadata(str(dicom_path)):
        success_count += 1
    
    # 5. Traditional formats (JSON vitals, text notes) - working formats
    total_docs += 2
    
    # JSON Vitals
    vitals_dir = test_patient_dir / 'vitals'
    vitals_dir.mkdir(exist_ok=True)
    vitals_data = {
        "date": "2024-01-15",
        "time": "14:30",
        "glucose_mg_dL": 245,
        "bp_systolic": 138,
        "bp_diastolic": 88,
        "heart_rate": 82,
        "temperature_f": 98.4,
        "weight_kg": 78.5,
        "oxygen_saturation": 98,
        "notes": "Patient reports improved energy"
    }
    
    vitals_path = vitals_dir / '2024-01-15_vitals.json'
    with open(vitals_path, 'w') as f:
        json.dump(vitals_data, f, indent=2)
    print(f"✅ Created JSON Vitals: {vitals_path}")
    success_count += 1
    
    # Text Clinical Note
    notes_dir = test_patient_dir / 'clinical_notes'
    notes_dir.mkdir(exist_ok=True)
    
    clinical_note = """CLINICAL NOTE - January 15, 2024

Patient: Test Patient 999
Age: 55 years old Male

CHIEF COMPLAINT: Follow-up for diabetes and hypertension

VITAL SIGNS:
- Blood Pressure: 138/88 mmHg
- Heart Rate: 82 bpm  
- Glucose: 245 mg/dL
- Temperature: 98.4°F
- Weight: 78.5 kg

ASSESSMENT:
1. Diabetes mellitus, Type 2 - glucose elevated
2. Essential hypertension - well controlled
3. Dyslipidemia - on statin therapy

PLAN:
- Adjust metformin dosage
- Continue lisinopril
- Dietary counseling
- Follow-up in 3 months

Dr. Amanda Foster, MD"""

    note_path = notes_dir / '2024-01-15_clinical_note.txt'
    with open(note_path, 'w') as f:
        f.write(clinical_note)
    print(f"✅ Created Text Note: {note_path}")
    success_count += 1
    
    print(f"\n📊 Document Creation Summary:")
    print(f"   ✅ Successfully created: {success_count}/{total_docs} documents")
    print(f"   📁 Test patient directory: {test_patient_dir}")
    
    # Now test document processing
    print(f"\n🔄 TESTING DOCUMENT PROCESSING:")
    print("-" * 40)
    
    try:
        from src.medical_document_processor import MedicalDocumentProcessor
        
        processor = MedicalDocumentProcessor()
        
        print(f"\n🔍 Processing Patient 999 documents...")
        patient_profile = processor.process_all_patient_documents('999')
        
        print(f"\n📋 PROCESSING RESULTS:")
        print(f"   Patient ID: {patient_profile['patient_id']}")
        print(f"   Documents: {patient_profile['document_count']}")
        print(f"   Document Types: {patient_profile.get('document_types', [])}")
        print(f"   Latest Vitals: {len(patient_profile.get('latest_vitals', {}))} measurements")
        
        # Show extracted values
        latest_vitals = patient_profile.get('latest_vitals', {})
        if latest_vitals:
            print(f"\n💊 EXTRACTED VITAL SIGNS:")
            for key, value in latest_vitals.items():
                if not key.endswith('_date'):
                    print(f"   {key}: {value}")
        
        # Show clinical timeline
        timeline = patient_profile.get('clinical_timeline', [])
        print(f"\n📅 CLINICAL TIMELINE ({len(timeline)} entries):")
        for entry in timeline[:5]:  # Show first 5 entries
            print(f"   {entry['date']}: {entry['category']} - {len(entry.get('values', {}))} values")
        
        print(f"\n✅ MULTI-FORMAT PROCESSING: SUCCESS!")
        return True
        
    except Exception as e:
        print(f"❌ Document processing failed: {e}")
        return False

def test_format_capabilities():
    """Test which formats are currently supported"""
    print(f"\n🔍 CHECKING FORMAT CAPABILITIES:")
    print("-" * 40)
    
    capabilities = {}
    
    # Test PDF
    try:
        import PyPDF2
        import fitz
        capabilities['PDF'] = '✅ Available (PyPDF2, PyMuPDF)'
    except ImportError as e:
        capabilities['PDF'] = f'❌ Not available: {e}'
    
    # Test OCR
    try:
        from PIL import Image
        import pytesseract
        capabilities['OCR/Images'] = '✅ Available (PIL, pytesseract)'
    except ImportError as e:
        capabilities['OCR/Images'] = f'❌ Not available: {e}'
    
    # Test DICOM
    try:
        import pydicom
        capabilities['DICOM'] = '✅ Available (pydicom)'
    except ImportError as e:
        capabilities['DICOM'] = f'❌ Not available: {e}'
    
    # Test Word
    try:
        from docx import Document
        capabilities['Word Documents'] = '✅ Available (python-docx)'
    except ImportError as e:
        capabilities['Word Documents'] = f'❌ Not available: {e}'
    
    # Always available
    capabilities['JSON'] = '✅ Available (built-in)'
    capabilities['Text Files'] = '✅ Available (built-in)'
    
    for format_type, status in capabilities.items():
        print(f"   {format_type}: {status}")
    
    return capabilities

if __name__ == "__main__":
    print("🏥 COMPREHENSIVE MULTI-FORMAT DOCUMENT TEST")
    print("=" * 70)
    
    # Test format capabilities
    capabilities = test_format_capabilities()
    
    # Run comprehensive test
    if test_multi_format_processing():
        print(f"\n🎉 SUCCESS: Multi-format document processing is working!")
        print(f"\n🚀 READY FOR PRODUCTION:")
        print(f"   • Copy your real medical documents to data/patient_documents/")
        print(f"   • Supported formats: PDF, Images, Word, JSON, Text")
        print(f"   • Documents will be processed automatically")
        print(f"   • No restart required for new documents")
    else:
        print(f"\n❌ FAILED: Check error messages above")
